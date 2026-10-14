"""Manual LLM benchmark without an API: an Excel answer sheet to fill in by hand from a free chat app.

  python -m benchmark.answer_sheet --make      writes benchmark/LLM_answer_sheet.xlsx
  (paste each prompt into a NEW or TEMPORARY chat, type the number it answers into the Answer column, fill in the
   LLM name, model and date on the Instructions sheet)
  python -m benchmark.answer_sheet --import    copies the answers into benchmark/responses.csv
  python -m benchmark.score                    scores them

Workload: run 1 covers every held-out profile (a fair comparison with the model on the same graduates); runs 2 and 3
cover a fixed random subset (default 30), enough to measure how consistent the answers are.
Existing answers in responses.csv are never overwritten.
"""
from __future__ import annotations

import argparse
import datetime as dt
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from benchmark.export_prompts import BENCH, DEFAULT_VARIANT, assert_fresh

SHEET = BENCH / "LLM_answer_sheet.xlsx"
REPEAT_SUBSET = 30
SEED = 20261013
INSTRUCTIONS = [
    "How to fill in this sheet",
    "1. Open the free chat app (for example ChatGPT). Turn on Temporary Chat, or start a NEW chat for every prompt, so "
    "earlier answers cannot influence later ones.",
    "2. On the Answers sheet, copy the whole Prompt cell, paste it into the chat and send it.",
    "3. Type only the number of months it answers into the Answer column. If it refuses or gives no number, copy "
    "what it said instead - the scorer records it as unparseable.",
    "4. Do not change the prompt, and do not answer the same row twice.",
    "5. Fill in the three cells below, then run:  python -m benchmark.answer_sheet --import",
    "",
    "Before you start: check that your consent wording or supervisor allows anonymised profiles to be entered into "
    "the chat app.",
]
DETAIL_CELLS = {"LLM name (e.g. ChatGPT)": "B11", "Model shown in the app": "B12", "Date(s) answered": "B13"}


def plan(prompts: pd.DataFrame, repeat: int = REPEAT_SUBSET, seed: int = SEED) -> pd.DataFrame:
    """Rows to answer: run 1 for every profile, runs 2-3 for a fixed random subset."""
    p = prompts[prompts["variant"] == DEFAULT_VARIANT].reset_index(drop=True)
    subset = set(np.random.default_rng(seed).choice(p["profile_id"], size=min(repeat, len(p)), replace=False))
    rows = [(pid, 1, txt) for pid, txt in zip(p["profile_id"], p["prompt_text"])]
    rows += [(pid, r, txt) for r in (2, 3) for pid, txt in zip(p["profile_id"], p["prompt_text"]) if pid in subset]
    out = pd.DataFrame(rows, columns=["profile_id", "run_number", "prompt_text"])
    out.insert(0, "No", range(1, len(out) + 1))
    return out


def make(bench: Path = BENCH, sheet: Path = SHEET, repeat: int = REPEAT_SUBSET) -> int:
    assert_fresh(bench)
    rows = plan(pd.read_csv(bench / "prompts.csv"), repeat)
    wb = openpyxl.Workbook()
    ins = wb.active
    ins.title = "Instructions"
    for i, line in enumerate(INSTRUCTIONS, start=1):
        ins.cell(i, 1, line).alignment = Alignment(wrap_text=True, vertical="top")
    ins["A1"].font = Font(bold=True, size=13)
    for label, ref in DETAIL_CELLS.items():
        ins[ref.replace("B", "A")] = label
        ins[ref.replace("B", "A")].font = Font(bold=True)
        ins[ref].fill = PatternFill("solid", fgColor="FFF2CC")
    ins.column_dimensions["A"].width = 110
    ins.column_dimensions["B"].width = 40
    ws = wb.create_sheet("Answers")
    ws.append(["No", "Profile", "Run", "Prompt (copy this whole cell)", "Answer (months)"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows.itertuples(index=False):
        ws.append([r.No, r.profile_id, r.run_number, r.prompt_text, None])
        ws.cell(ws.max_row, 4).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(ws.max_row, 5).fill = PatternFill("solid", fgColor="FFF2CC")
    for col, w in zip("ABCDE", (6, 9, 6, 100, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(sheet)
    return len(rows)


def import_answers(bench: Path = BENCH, sheet: Path = SHEET) -> dict:
    assert_fresh(bench)
    wb = openpyxl.load_workbook(sheet, data_only=True)
    details = {label: wb["Instructions"][ref].value for label, ref in DETAIL_CELLS.items()}
    if not details["LLM name (e.g. ChatGPT)"]:
        raise SystemExit("Fill in the LLM name on the Instructions sheet first.")
    answers = pd.DataFrame(list(wb["Answers"].iter_rows(min_row=2, values_only=True)),
                           columns=["No", "profile_id", "run_number", "prompt_text", "answer"])
    answers = answers[answers["answer"].notna() & (answers["answer"].astype(str).str.strip() != "")]
    path = bench / "responses.csv"
    resp = pd.read_csv(path, dtype={"raw_response": "string", "llm_name": "string", "model_version": "string",
                                    "date": "string"})
    filled = skipped = 0
    for a in answers.itertuples(index=False):
        m = (resp["profile_id"] == a.profile_id) & (resp["run_number"] == int(a.run_number)) & \
            (resp["variant"] == DEFAULT_VARIANT)
        if not m.any():
            continue
        i = resp.index[m][0]
        if pd.notna(resp.at[i, "raw_response"]) and str(resp.at[i, "raw_response"]).strip():
            skipped += 1                                  # never overwrite a recorded answer
            continue
        resp.at[i, "raw_response"] = str(a.answer).strip()
        resp.at[i, "llm_name"] = str(details["LLM name (e.g. ChatGPT)"])
        resp.at[i, "model_version"] = str(details["Model shown in the app"] or "")
        resp.at[i, "date"] = str(details["Date(s) answered"] or dt.date.today().isoformat())
        filled += 1
    resp.to_csv(path, index=False)
    return {"answers_in_sheet": len(answers), "copied": filled, "already_recorded_kept": skipped}


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--make", action="store_true")
    g.add_argument("--import", dest="do_import", action="store_true")
    ap.add_argument("--repeat", type=int, default=REPEAT_SUBSET, help="profiles answered three times (default 30)")
    a = ap.parse_args()
    if a.make:
        print(f"wrote {SHEET} with {make(repeat=a.repeat)} prompts to answer")
    else:
        print(import_answers())
